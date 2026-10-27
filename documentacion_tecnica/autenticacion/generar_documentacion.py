"""Genera los entregables del sprint desde evidencia local, sin credenciales ni llamadas autenticadas."""
from pathlib import Path
import json
import re
import hashlib
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from lxml import etree

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
EV = BASE / 'evidencias'
DATE = '27 de octubre de 2026'
META = json.loads((EV / 'contexto-revision.json').read_text(encoding='utf-8'))
SYMBOLS = json.loads((EV / 'inventario-simbolos.json').read_text(encoding='utf-8'))
OK = 'Implementado y validado'
PART = 'Parcialmente implementado'
PEND = 'Pendiente'
NV = 'No verificable'
OLIVE = '78813E'
DARK = '172419'

def readlog(name):
    raw = (EV / name).read_bytes()
    return raw.decode('utf-16') if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else raw.decode('utf-8-sig', errors='replace')

focused = readlog('pruebas-autenticacion.log')
full = readlog('pruebas-completas.log')
assert re.search(r'Tests\s+53 passed', focused), 'Actualizar resultados si cambia la ejecución enfocada'
assert re.search(r'Tests\s+10 failed\s*\|\s*166 passed', full), 'Actualizar resultados globales'

# Cada registro corresponde a un archivo existente; no se copian cuerpos con credenciales.
inventory = [
('src/app/pages/login/login.component.ts', 'LoginComponent',
 'Formulario reactivo de ingreso, estado de envío y mensajes seguros.',
 'FormBuilder, AuthService, Router, DestroyRef; plantilla y fondo TopographyBackground.',
 'login(usuario,password) a través de AuthService; no HTTP directo.',
 'Evita envíos duplicados; takeUntilDestroyed cancela al salir; finalize libera el botón.',
 '401/400: credenciales; 422: revisar datos; red/servidor: mensaje controlado; contrato o timeout: error genérico.',
 'El diseño no cambia. recordarUsuario no implementa persistencia de sesión. La contraseña existe en el control del formulario durante su vida; no se almacena como dato de sesión.'),
('src/app/core/auth/auth.service.ts', 'AuthService',
 'Fuente central de usuario, access token, expiración, refresh reservado, logout y generación local de sesión.',
 'HttpClient, ApiConfigService, Router, CatalogosService, RxJS; guards, interceptores, Sidebar y LoginComponent.',
 'POST ApiConfigService.loginUrl, actualmente /api/auth/token. Ninguna solicitud de refresh, logout o perfil.',
 'Tokens en memoria; validación completa previa; copia de usuario/catálogos; cancelación de login anterior; limpieza de claves antiguas; respuestas aisladas por generación.',
 'Timeout de 30000 ms; error de login limpia estado sin cancelar su propia notificación. Fallos de limpieza/navegación se advierten y no restauran la sesión.',
 'isAuthenticated revisa la fecha absoluta derivada de expira_en; no decodifica JWT. getCurrentUser/getRole no comprueban por sí mismos expiración. Mock preexistente solo si useMockAuth y no producción, actualmente desactivado. No se reproducen sus credenciales en este documento.'),
('src/app/core/auth/auth.models.ts', 'Interfaces de autenticación',
 'Define LoginResponse, AuthSession, UsuarioAutenticado/AuthUser, LoginRequest y los cinco catálogos.',
 'AuthService, validador, CatalogosService y pruebas.',
 'Modelo del bootstrap consumido por login; no ejecuta HTTP.',
 'Estructuras conocidas sin any; campos del backend conservados.',
 'Las interfaces no validan JSON en ejecución: lo hace validarLogin.',
 'El contrato activo exige usuario.permisos:string[], rol:string y expira_en:number. No existe JwtPayload activo. Los datos de tokens y usuario son campos separados, aunque todavía se exponen signals escribibles.'),
('src/app/core/auth/auth-response.validator.ts', 'validarLogin',
 'Acota una entrada no confiable a LoginResponse verificando todos los elementos del bootstrap.',
 'LoginResponse; invocado por mapLoginResponse antes de publicar estado.',
 'Sin llamadas; procesa respuesta de login.',
 'Tokens no vacíos, bearer, expiración positiva y finita; usuario y cinco arreglos tipados; IDs de catálogo no duplicados.',
 'Lanza Error controlado ante datos faltantes o incompatibles; AuthService limpia la sesión.',
 'No verifica firma, exp, sub, rol numérico ni revocación. No limita tamaño de arrays ni valida todas las reglas semánticas de datos de negocio.'),
('src/app/core/auth/catalogos.service.ts', 'CatalogosService y NodoUnidad',
 'Conserva bootstrap de catálogos y ofrece consultas de jerarquía.',
 'CatalogosLogin, UnidadMilitar, signals; AuthService carga y limpia.',
 'Sin HTTP; no requiere endpoint adicional.',
 'Copia al cargar y estado público de solo lectura a nivel de signal; ciclos detenidos mediante Set.',
 'Búsquedas sin coincidencia devuelven undefined; rutas inexistentes devuelven arreglo vacío. structuredClone puede lanzar.',
 'Usa unidad_padre_id para raíces, hijas, superior, árbol y ruta. Las huérfanas aparecen como raíces del árbol. Los objetos devueltos no están profundamente congelados.'),
('src/app/core/auth/authorization-policy.ts', 'Política numérica preparada',
 'Reglas puras para ContextoAutorizacion con rol ADMINISTRADOR/USUARIO y permisos number[].',
 'Pruebas authorization-policy.spec.ts; sin dependencias Angular.',
 'Sin HTTP ni conexión a JWT.',
 'Administrador sin lista permite IDs válidos; demás casos exigen asignación; sesión ausente, rol inválido e IDs no positivos/enteros se deniegan.',
 'Devuelve false, sin lanzar, para contextos inválidos; listas vacías de requisitos no autorizan.',
 'NO se importa desde AuthService, rutas o menús. Las pruebas prueban la política aislada, no autorización operacional. No confundir con tienePermiso(codigo:string).'),
('src/app/core/auth/auth.guard.ts', 'authGuard',
 'Protege las rutas hijas del AppShell por presencia de sesión local vigente.',
 'AuthService, Router, APP_PATHS; app.routes.ts.',
 'Sin HTTP.',
 'Retorna UrlTree a /login en lugar de navegar imperativamente.',
 'No hay sesión: redirección; token vencido: AuthService limpia estado.',
 'No comprueba rol ni permisos de una funcionalidad; no valida vigencia remota.'),
('src/app/core/auth/guest.guard.ts', 'guestGuard',
 'Reserva login a visitantes sin sesión local activa.',
 'AuthService, Router, APP_PATHS.',
 'Sin HTTP.',
 'Con sesión retorna UrlTree a /home.',
 'Sin sesión permite acceso; la expiración se evalúa en AuthService.',
 'No es un guard de administrador ni sustituye autorización del servidor.'),
('src/app/core/auth/auth-http-context.ts', 'Contextos HTTP',
 'Declara SKIP_AUTH_INTERCEPTOR y SKIP_ERROR_INTERCEPTOR con valor por defecto false.',
 'AuthService e interceptores auth/error.',
 'Login utiliza ambos contextos de exclusión.',
 'Evita adjuntar credenciales previas y evita logout global ante credenciales incorrectas.',
 'No procesa errores; controla qué interceptor los atiende.',
 'No existe contexto de reintento/refresh en esta versión.'),
('src/app/core/interceptors/auth.interceptor.ts', 'authInterceptor',
 'Adjunta Bearer a backends configurados, controla cache y filtra éxitos obsoletos.',
 'ApiConfigService, AuthService, SKIP_AUTH_INTERCEPTOR, RxJS filter.',
 'Actúa en las bases configuradas; excluye token Bearer del login.',
 'Comparación de origen y límite de ruta; cache:no-store en backends; proveedores externos no reciben credenciales.',
 'No renueva ni reintenta; deja errores al siguiente interceptor. Solo filtra respuestas exitosas por generación.',
 'X-API-Key se conserva como compatibilidad pero su signal queda null con el contrato vigente. No envía refresh token.'),
('src/app/core/interceptors/error.interceptor.ts', 'errorInterceptor',
 'Centraliza 401/403/5xx de los backends configurados.',
 'ApiConfigService, AuthService, SKIP_ERROR_INTERCEPTOR.',
 'No crea solicitudes propias salvo logout local por AuthService.',
 '401 protegido solo afecta a la generación actual; login/exteriores excluidos; 403 no cierra sesión.',
 'Rethrow de errores; 403 advierte; 5xx registra error completo en consola.',
 'Hallazgo: HttpErrorResponse completo puede contener detalles sensibles del backend. Redactar logs en futuro sprint; no se cambia el comportamiento en esta fase.'),
('src/app/core/interceptors/loader.interceptor.ts', 'loaderInterceptor',
 'Activa LoaderService únicamente para solicitudes con LOADER_OPERATION.',
 'LoaderService; defer/finalize.',
 'Transversal; no define endpoints.',
 'No clasifica peticiones automáticas como actividad humana ni implementa temporizador de sesión.',
 'finalize libera operación también en error/cancelación.',
 'No sustituye el signal isSubmitting del login; incluido por su orden en la cadena global.'),
('src/app/core/services/api-config.service.ts', 'ApiConfigService',
 'Centraliza bases, URLs de login/logout/refresh y límites de destino.',
 'environment.api y environment.endpoints; AuthService e interceptores.',
 'loginUrl=/api/auth/token; logoutUrl=/api/auth/logout y refreshTokenUrl=/api/auth/refresh configuradas pero no consumidas.',
 'matchesBaseUrl verifica origin y frontera de pathname, evitando prefijos parecidos.',
 'URL inválida produce false; base faltante utiliza respaldo auth.',
 'joinUrl concatena, no descubre contrato ni convierte las rutas existentes a v1. La base de fotos se resuelve desde el origen auth.'),
('src/app/core/services/navigation.service.ts', 'NavigationService',
 'Preparado para registrar NavigationEnd y nombre de usuario en un signal en memoria.',
 'Router, AuthService, DestroyRef; RxJS y takeUntilDestroyed.',
 'Sin HTTP.',
 'No guarda tokens ni credenciales en sus entradas.',
 'No constituye auditoría persistente ni demuestra eventos del servidor.',
 'No se encontró inyección/instanciación en app.ts ni otros consumidores; no afirmar que el registro esté activo. Si se activa, definir retención y limpieza al cambiar de usuario.'),
('src/app/pages/sidebar/sidebar.ts', 'Sidebar y MenuItem',
 'Navegación de escritorio/móvil, visibilidad estática, cierre de menús y acción Salir.',
 'AuthService, ElementRef, RouterLink, RouterLinkActive, APP_PATHS.',
 'Logout delegado a AuthService, actualmente solo local.',
 'Salir es botón; cierra menús y utiliza cierre centralizado.',
 'No hace HTTP ni asigna permisos ante errores.',
 'canShowItem solo evalúa item.permiso !== false; permisoCodigo es metadato no aplicado. El menú no filtra por IDs del backend.'),
('src/app/layout/app-shell.ts', 'AppShell',
 'Contenedor protegido y propietario del ciclo de vida de vistas/caché privada.',
 'AuthService, MapaElementosService, IconosService, MeasurementService, MapActionsService; PuntosDemoService por instancia.',
 'No crea endpoints de autenticación.',
 'Registra callback de limpieza; destruye vistas al navegar fuera y evita heredar caché de otro usuario.',
 'AuthService aísla errores de callbacks; ngOnDestroy retira registro y limpia estado.',
 'La documentación no altera herramientas geográficas ni sus algoritmos.'),
('src/app/app.routes.ts', 'routes',
 'Define login con guestGuard y AppShell con canActivateChild authGuard.',
 'Componentes cargados de forma diferida y APP_ROUTES.',
 'Rutas frontend; no son endpoints del backend.',
 'Rutas Home, Ingreso, Permisos, Vulcano y VulcanoLog requieren autenticación.',
 'Fallback a Home; guards generan redirección a login cuando corresponde.',
 'La ruta diagnóstica de terreno está fuera del shell. No hay guard por IDs ni módulo de seguridad administrador.'),
('src/app/app.config.ts', 'appConfig',
 'Registra Router y HttpClient con interceptores funcionales.',
 'routes, loaderInterceptor, authInterceptor y errorInterceptor.',
 'Sin rutas HTTP nuevas.',
 'Orden de salida: loader → auth → error; las respuestas recorren la cadena inversa.',
 'Registra escuchas de errores globales Angular.',
 'No incorpora inicializador de recuperación de sesión ni /yo.'),
('src/environments/environment.shared.ts', 'sharedEnvironment',
 'Define timeout, bases de servicios, endpoints y useMockAuth=false.',
 'environment.ts y environment.production.ts lo extienden.',
 'Base auth /api/auth; sufijos /token, /logout, /refresh.',
 'No contiene secreto de firma JWT. El modo mock requiere además no producción.',
 'No verifica conectividad ni despliegue.',
 'Otros servicios conservan direcciones localhost: verificar configuración de producción y transporte HTTPS en despliegue real, sin inferir que ya estén expuestos.'),
('proxy.conf.cjs', 'Proxy de desarrollo',
 'Reenvía /api/auth/** al backend local y elimina el prefijo /api/auth.',
 'ng serve; backend de desarrollo.',
 '/api/auth/token se traduce a /token en http://127.0.0.1:8010.',
 'Configuración de desarrollo, no evidencia del proxy o HTTPS en producción.',
 'Errores de conexión llegan al HttpClient/formulario.',
 'No configura /api/v1/auth/**. Los contratos v1 requieren verificar también el proxy, no solo cambiar strings.'),
('src/app/core/constants/storage-keys.constants.ts', 'STORAGE_KEYS',
 'Enumera claves propias antiguas de token, usuario, clave API y expiración.',
 'AuthService y pruebas de aislamiento.',
 'Sin HTTP.',
 'Se usan para retirar solo datos de autenticación propios en ambos almacenamientos.',
 'AuthService captura errores de acceso al storage.',
 'La existencia de claves no significa que se escriban credenciales: la sesión actual solo vive en memoria.'),
('src/app/core/constants/routes.constants.ts', 'APP_ROUTES y APP_PATHS',
 'Nombres relativos y absolutos de navegación.',
 'Guards, Router, Sidebar, LoginComponent y AuthService.',
 'No define API REST.',
 'Mantiene las redirecciones centralizadas.',
 'No maneja errores de red.',
 'No documenta equivalencias entre funciones e IDs de permisos.'),
('src/app/pages/login/login.component.html', 'Vista de login',
 'Presenta controles, envío, errores y contenidos institucionales existentes.',
 'LoginComponent, formularios reactivos y TopographyBackground.',
 'Sin HTTP directo.',
 'Consume estados y validación del componente; no muestra tokens.',
 'Muestra loginError controlado por el componente.',
 'No se modifica el diseño en este trabajo; CSS se revisa como referencia de paleta.'),
('src/app/pages/sidebar/sidebar.html', 'Vista de navegación',
 'Muestra grado, nombre completo, rol y acciones de salida.',
 'Sidebar y AuthService.',
 'Sin HTTP directo.',
 'No imprime tokens; navegación y salida delegadas.',
 'Sin procesamiento de respuestas.',
 'Información de identidad visible, sin implementar una vista /yo ni panel de seguridad.'),
]
for row in inventory:
    assert (ROOT / row[0]).exists(), row[0]

endpoints = [
('/api/auth/token → /token', 'POST confirmado', 'username/password URL-encoded; respuesta esperada LoginResponse: access_token, refresh_token, tipo_token, expira_en, usuario, catalogos.',
 'Sin Bearer de sesión anterior; contextos excluyen auth/error.', 'AuthService.loginWithBackend → LoginComponent',
 'OpenAPI local declara 200 (esquema libre) y 422. 401/400, red y 5xx cubiertos por manejo/pruebas HTTP simuladas, no verificados en servidor.',
 PART, 'Código y OpenAPI local; pruebas con HttpTestingController. No login real ejecutado.'),
('/api/v1/auth/login', 'No confirmado', 'Bootstrap solicitado; método, cuerpo y esquema real pendientes.', 'No confirmada', 'Ninguno para v1', 'No confirmados', PEND, 'Ruta en requerimiento; ausente del OpenAPI local consultado.'),
('/api/v1/auth/refresh', 'No confirmado', 'Transporte del refresh y respuesta de rotación pendientes.', 'No confirmada', 'Ninguno', 'No confirmados', PEND, 'Refresh conservado en memoria, sin llamada ni operación compartida.'),
('/api/v1/auth/logout', 'No confirmado', 'Cuerpo y confirmación de revocación pendientes.', 'Access token vigente solicitado; contrato no verificado', 'Ninguno remoto; logout actual es local', 'No confirmados', PEND, 'No hay HttpClient hacia logout en AuthService.'),
('/api/v1/auth/yo', 'No confirmado', 'Esquema del perfil y recuperación pendientes.', 'No confirmada', 'Ninguno', 'No confirmados', PEND, 'Sin recuperación de sesión persistida.'),
('/api/v1/auth/seguridad', 'No confirmado', 'Estructura de intentos/IPs/Redis pendiente.', 'Administrador requerido por especificación, no integración comprobada', 'Ninguno', 'No confirmados', PEND, 'Sin componente ni guard de administrador.'),
('/api/v1/auth/seguridad/bloqueo/{ip}', 'No confirmado', 'IP, cuerpo, método y salida pendientes; confirmación manual requerida.', 'Administrador requerido por especificación', 'Ninguno', 'No confirmados', PEND, 'No se enviaron solicitudes de desbloqueo.'),
]
histories = [
('HU-01', 'Como usuario necesito iniciar sesión con bootstrap válido.', 'Formulario URL-encoded; solo autenticar con tokens/usuario/catálogos completos; credenciales incorrectas dejan formulario utilizable.', OK, 'login.component.spec.ts; login-contract.spec.ts; api-connection.spec.ts. Validación con backend simulado.'),
('HU-02', 'Como usuario necesito renovar la sesión mientras mantengo actividad.', 'Access 25 min, ventana 60 min, rotación única compartida y sin extender sesión por telemetría.', PEND, 'No hay refresh, listeners de actividad ni temporizador; 25/60 min son requisitos, no valores operativos verificados.'),
('HU-03', 'Como administrador necesito limitar funciones mediante roles y permisos.', 'ADMINISTRADOR vacío permite todo; otras listas limitan; aplicar a menú, ruta y acción con IDs documentados.', PART, 'authorization-policy.spec.ts valida funciones puras; no conexión operacional. AuthService sigue usando string[].'),
('HU-04', 'Como usuario necesito cerrar mi sesión e impedir reutilización local.', 'Limpiar tokens/catálogos, cancelar login pendiente, descartar respuestas anteriores y destruir vistas privadas; revocación remota diferenciada.', PART, 'logout.spec.ts y app-shell.spec.ts validan cierre local. Redis/revocación pendientes.'),
('HU-05', 'Como administrador necesito consultar y desbloquear IPs autorizadas.', 'Solo administrador; contrato confirmado; mostrar IP; confirmación manual y refrescar lista tras éxito.', PEND, 'Sin módulo ni consumidor HTTP de seguridad.'),
('HU-06', 'Como desarrollador necesito catálogos de sesión sin consultas duplicadas.', 'Cinco catálogos cargados desde bootstrap; jerarquía por unidad_padre_id y limpieza al salir.', OK, 'login-contract.spec.ts: cinco catálogos, búsquedas, árbol, ruta y corte de ciclos.'),
('HU-07', 'Como usuario necesito recuperar un perfil válido al recargar.', 'Persistencia respaldada por backend, consulta /yo y recuperación sin confiar solo en JWT almacenado.', PEND, 'Estado en memoria; recarga requiere nuevo login. No /yo ni persistencia segura confirmada.'),
('HU-08', 'Como mantenedor necesito documentación verificable del sprint.', 'Inventario, trazabilidad, evidencias de ejecución, TSDoc sin cambios funcionales y Word legible.', OK, 'Entregables presentes; igualdad del JavaScript transpilado y validación estructural del DOCX.'),
]
tests = [
('PV-01', 'Contrato bootstrap y atomicidad', 'Rechazar JSON incompleto y cargar usuario/catálogos válidos', 'Selección enfocada aprobada; login-contract.spec.ts y api-connection.spec.ts', OK, 'pruebas-autenticacion.log'),
('PV-02', 'Login y mensajes', 'Evitar duplicados; manejar 401/422/red/500 sin detalles sensibles', 'Casos del componente aprobados con backend simulado', OK, 'login.component.spec.ts; pruebas-autenticacion.log'),
('PV-03', 'Headers y límites de backend', 'Bearer solo a destinos configurados; no refresh en solicitudes ordinarias', 'Casos API aprobados; no petición de refresh', OK, 'api-connection.spec.ts; pruebas-autenticacion.log'),
('PV-04', 'Guards y logout', 'Bloquear rutas tras salir, descartar respuestas antiguas, limpiar vistas privadas', 'Casos de logout/AppShell aprobados', OK, 'logout.spec.ts; app-shell.spec.ts; pruebas-autenticacion.log'),
('PV-05', 'Reglas numéricas', 'Aplicar matriz de roles y denegar IDs/contextos inválidos', 'Política pura aprobada; no demuestra protección de módulos', OK, 'authorization-policy.spec.ts; pruebas-autenticacion.log'),
('PV-06', 'Expiración y almacenamiento', 'Caducar según expira_en y no restaurar credenciales locales', 'Casos con reloj simulado aprobados', OK, 'api-connection.spec.ts; login-contract.spec.ts'),
('PV-07', 'Suite completa del repositorio', 'Todas las pruebas del estado actual deberían aprobar', '166 aprobadas, 10 fallidas; 32 archivos aprobados y 5 fallidos, total 176/37', PART, 'pruebas-completas.log'),
('PV-08', 'Build desarrollo', 'Compilar TypeScript y templates Angular', 'Comando terminó correctamente', OK, 'compilacion-desarrollo.log'),
('PV-09', 'Build producción', 'Cumplir compilación y presupuestos', 'Bloqueado por CSS Hoja Milimetrada: 19,26 kB > 18 kB', PART, 'compilacion-produccion.log'),
('PV-10', 'Ausencia de cambio funcional en TSDoc', 'Conservar JavaScript transpilado en 15 archivos revisados', '15/15 hashes idénticos con removeComments=true', OK, 'verificacion-solo-comentarios.json'),
('PV-11', 'Login/refresh/logout/seguridad reales v1', 'Contrato y comportamiento real de FastAPI/Redis', 'No ejecutados; endpoints ausentes del OpenAPI local', NV, 'contexto-revision.json'),
('PV-12', 'Word', 'Paquete OOXML íntegro y apertura sin reparación', 'Validación estructural automática; resultado de apertura Word registrado aparte al finalizar', 'Implementado sin validar', 'validacion-documento.json (generado tras abrir el documento)'),
]
validation_path = EV / 'validacion-documento.json'
if validation_path.exists():
    validation = json.loads(validation_path.read_text(encoding='utf-8-sig'))
    if validation.get('apertura_microsoft_word'):
        tests[-1] = ('PV-12', 'Word', 'Paquete OOXML íntegro y apertura sin reparación',
                     'Abierto con Microsoft Word mediante COM; campos/páginas actualizados y PDF exportado; sin aprobación editorial',
                     OK, 'validacion-documento.json; PDF de presentación')
risks = [
('R-01', 'Alta', 'Autorización operacional por IDs sin integrar: guards solo autentican y menús usan booleanos estáticos.', 'Un usuario autenticado podría visualizar/activar funciones sin filtrado local por permisos; backend debe impedir operaciones no autorizadas.', 'Obtener seed/contrato, conectar política numérica a rutas/menús/acciones y probar cada función. No asumir una vulnerabilidad del backend sin evidencia.'),
('R-02', 'Alta', 'Logout únicamente local, sin revocación confirmada.', 'No garantiza invalidación remota de un token previamente emitido.', 'Confirmar contrato logout y su semántica de revocación; diferenciar fallo remoto y cierre local.'),
('R-03', 'Alta', 'Sin refresh, control humano de actividad ni sesión deslizante.', 'Interrumpe continuidad al vencer access token; requisito de inactividad no cubierto.', 'Implementar rotación single-flight y ventana de actividad tras confirmar contrato; nunca renovar indefinidamente por tráfico automático.'),
('R-04', 'Media', 'Expiración basada en expira_en local, sin lectura del JWT real.', 'No detecta diferencias con exp o revocación antes de recibir rechazo del backend.', 'Confirmar claims y comparar vigencias sin intentar verificar firma con secretos en Angular; backend mantiene autoridad.'),
('R-05', 'Media', 'errorInterceptor registra HttpErrorResponse completo para 5xx.', 'El cuerpo del servidor podría incluir datos sensibles en consola.', 'Redactar detalle, limitar logs por entorno y definir telemetría sin credenciales en un cambio funcional separado.'),
('R-06', 'Media', 'Signals de AuthService escribibles desde consumidores; modelos/catálogos no profundamente inmutables.', 'Posibles inconsistencias de estado por escrituras externas o mutación accidental.', 'Encapsular estado con signals privados/asReadonly y devolver vistas inmutables en sprint posterior.'),
('R-07', 'Media', 'Recarga sin persistencia ni /yo.', 'Se pierde sesión en memoria al recargar; no hay recuperación confirmada.', 'Definir mecanismo de persistencia respaldado por backend; no persistir refresh por conveniencia.'),
('R-08', 'Media', 'Transporte y bases de producción requieren revisión de despliegue.', 'Direcciones localhost en servicios auxiliares no prueban HTTPS ni enrutamiento institucional.', 'Verificar Nginx/HTTPS, CORS y URLs por ambiente con equipo de backend; no inferir configuración del servidor a partir del proxy de desarrollo.'),
('R-09', 'Media', 'Suite global con fallos y build producción bloqueado por presupuesto CSS.', 'Impide afirmar que el producto completo está validado o listo para publicar.', 'Corregir en tareas separadas las pruebas/causas y revisar presupuesto de estilos sin alterar esta documentación funcionalmente.'),
('R-10', 'Baja', 'Documentación histórica diverge de código: README menciona arranque de NavigationService y contrato antiguo.', 'Confusión durante transferencia del conocimiento.', 'Usar este inventario como corte vigente y actualizar material histórico con aprobación de alcance. No atribuir auditoría activa no encontrada.'),
('R-11', 'Media', 'Ruta de diagnóstico de terreno fuera del guard de AppShell.', 'Superficie accesible sin autenticación por configuración actual.', 'Confirmar con responsables si debe ser pública; no cambiar rutas en esta fase documental.'),
]
trace = [
('HU-01', 'Login y bootstrap íntegro', 'LoginComponent; AuthService; validarLogin', '/api/auth/token → /token', 'PV-01, PV-02', OK),
('HU-01', 'Bearer y protección por sesión', 'authInterceptor; authGuard; AppShell', 'Backends configurados', 'PV-03, PV-04', OK),
('HU-02', 'Rotación e inactividad 25/60', 'AuthService (sin implementación)', '/api/v1/auth/refresh', 'PV-11', PEND),
('HU-03', 'Matriz numérica pura', 'authorization-policy.ts', 'Sin HTTP', 'PV-05', OK),
('HU-03', 'Permisos efectivos por función', 'Sidebar; app.routes.ts; AuthService', 'Contrato/seed sin confirmar', 'Sin prueba operacional por ID', PART),
('HU-04', 'Limpieza y aislamiento local', 'AuthService; AppShell; interceptores', 'Sin HTTP de logout', 'PV-04', OK),
('HU-04', 'Revocación remota', 'Sin consumidor implementado', '/api/v1/auth/logout', 'PV-11', PEND),
('HU-05', 'Consulta y desbloqueo IP', 'Sin módulo implementado', '/api/v1/auth/seguridad[/bloqueo/{ip}]', 'PV-11', PEND),
('HU-06', 'Catálogos y jerarquía', 'CatalogosService; auth.models.ts', 'Bootstrap del login', 'PV-01', OK),
('HU-07', 'Recuperación de perfil', 'Sin consumidor implementado', '/api/v1/auth/yo', 'PV-11', PEND),
('HU-08', 'Documentación sin alterar lógica', '13 TS comentados; 15 comparados; cinco entregables', 'OpenAPI solo lectura', 'PV-08, PV-10, PV-12', OK),
]

def mdtable(headers, rows):
    escape = lambda v: str(v).replace('|', '\\|').replace('\n', ' ')
    return '| ' + ' | '.join(headers) + ' |\n|' + '|'.join(['---']*len(headers)) + '|\n' + '\n'.join('| ' + ' | '.join(map(escape,r)) + ' |' for r in rows) + '\n'

architecture = '''# Arquitectura de autenticación — VULCANO

Fecha de corte: 27/10/2026. Versión documental 1.0. Revisión del árbol de trabajo, no certificación de producción.

## Estado comprobado

Angular standalone y signals centralizan la sesión en AuthService. LoginComponent envía username/password URL-encoded a /api/auth/token; proxy.conf.cjs elimina /api/auth y llama /token en el backend local. validarLogin valida usuario, tokens y cinco catálogos; CatalogosService copia los catálogos antes de publicar acceso. La expiración usa Date.now()+expira_en*1000. Access y refresh solo permanecen en memoria. Recargar requiere login; no hay /yo, refresh ni revocación remota implementados.

El login omite authInterceptor y errorInterceptor mediante HttpContext; el componente muestra errores controlados. La cadena global está registrada como loader → auth → error. El interceptor auth limita Bearer a bases configuradas, usa cache:no-store y descarta éxitos de sesiones previas. El interceptor error cierra localmente ante 401 protegido actual; 403 conserva sesión. Logout limpia estado, claves propias, cookies visibles conocidas y callbacks privados; no borra cookies HttpOnly del servidor ni revoca remotamente.

## Diagrama de arquitectura

```mermaid
flowchart TB
  UI[LoginComponent / Sidebar] --> AUTH[AuthService: signals y generación local]
  AUTH --> VALID[validarLogin y CatalogosService]
  AUTH --> HTTP[HttpClient / ApiConfigService]
  HTTP --> PROXY[Proxy desarrollo: /api/auth a raíz]
  PROXY --> API[FastAPI local: POST /token]
  AUTH --> GUARD[Guards y AppShell]
  GUARD --> VIEW[Vistas autenticadas]
  POLICY[Política numérica aislada, no integrada] -. pendiente .-> GUARD
  API -. no verificable .-> BACKEND[JWT / BD / Redis internos]
```

## Diagrama de secuencia

```mermaid
sequenceDiagram
  actor U as Usuario
  participant L as LoginComponent
  participant A as AuthService y catálogos
  participant F as FastAPI vía proxy
  U->>L: Envía formulario válido
  L->>A: login(usuario,password)
  A->>F: POST /token (URL-encoded)
  Note over F: Validación de credenciales y firma JWT: internos no auditados
  F-->>A: Bootstrap esperado por el frontend
  A->>A: Validar todo y copiar catálogos/usuario
  alt Respuesta válida
    A-->>L: next(LoginResponse)
    L->>L: Navegar /home
  else Error HTTP, timeout o contrato
    A->>A: Limpiar sesión
    A-->>L: error controlado
  end
  Note over A,F: Sin refresh, /yo o logout remoto implementados
```

## Autorización: dos estados distintos

AuthService opera con usuario.permisos:string[] y comparación exacta; no hay excepción activa de administrador. Sidebar utiliza booleanos estáticos y authGuard solo autentica. authorization-policy.ts implementa en aislamiento ADMINISTRADOR vacío=total, administrador con lista=limitado y USUARIO vacío=denegado. No conecta los IDs a funciones reales ni decodifica JWT. No existe JwtPayload activo ni lectura de sub/rol/permisos/exp/iat/jti desde un token. Estos claims son requisitos del contrato v1, aún no evidencia de su contenido real. La validación criptográfica y autorización definitiva pertenecen al backend.

## Sesión y persistencia

La ventana de 25 minutos de access y 60 minutos de inactividad, rotación de refresh, rechazo de renovación concurrente y actividad humana son requisitos pendientes. El código acepta expira_en positivo del bootstrap; no fija 25 minutos para el backend. Redis, base de datos de identidades, lista de revocación, fuerza bruta e IPs no están comprobados con código backend. La única evidencia consultada del servidor es OpenAPI, sin autenticarse; su respuesta 200 de /token no describe un esquema estricto. No se ejecutaron operaciones remotas de seguridad.

## Contratos y próximos pasos

Obtener OpenAPI v1, ejemplo anónimo de payload/bootstrap y seed de permisos. Confirmar métodos/cuerpos/respuestas, expiración, transporte del refresh y persistencia antes de integrar. Conectar permisos a guard, menú y acción; implementar refresh compartido con un solo reintento y actividad humana eficiente; añadir revocación confirmada, /yo y módulo administrador. Mantener catálogos fuera del JWT y reutilizar CatalogosService.
'''
(BASE / 'ARQUITECTURA_AUTENTICACION.md').write_text(architecture,encoding='utf-8')
(BASE / 'MATRIZ_TRAZABILIDAD_AUTENTICACION.md').write_text('# Matriz de trazabilidad de autenticación\n\n'+mdtable(['Historia','Requisito','Archivo','Endpoint','Prueba','Estado'],trace)+'\n## Historias y aceptación\n\n'+mdtable(['Historia','Necesidad','Criterios de aceptación','Estado','Evidencia'],histories)+'\n## Pruebas ejecutadas y pendientes\n\n'+mdtable(['ID','Caso','Esperado','Obtenido','Estado','Evidencia'],tests)+'\nEstados permitidos: Implementado y validado; Implementado sin validar; Parcialmente implementado; Pendiente; No verificable. La validación aislada de una política no implica integración a rutas. PV-12 se completa con el acta automática de apertura del Word.\n',encoding='utf-8')

invmd = '# Inventario de archivos de autenticación\n\nFecha: 27/10/2026. Rutas relativas a la raíz del proyecto. Las firmas provienen del AST TypeScript; no se incluyen cuerpos con credenciales.\n'
for i,row in enumerate(inventory,1):
    file,name,role,deps,ep,sec,err,obs = row
    invmd += f'\n## F-{i:02d} — {name}\n\nRuta: `{file}`\n\nResponsabilidad: {role}\n\nDependencias/consumidores: {deps}\n\nEndpoints: {ep}\n\nSeguridad: {sec}\n\nErrores: {err}\n\nObservaciones: {obs}\n'
    sigs = [s for s in SYMBOLS if s['archivo']==file]
    if sigs:
        invmd += '\n'+mdtable(['Método','Entrada/parámetros','Retorno','Línea'],[(s['nombre'],s['entrada'] or 'Sin parámetros',s['retorno'],s['linea']) for s in sigs])
    else:
        invmd += '\nFunciones/entradas/retornos: constantes, interfaces, plantilla o función interceptor/guard descrita arriba; sin métodos de clase adicionales inventados. Guards retornan boolean/UrlTree; interceptores reciben solicitud y next y retornan flujo HTTP.\n'
invmd += '\n## Archivos de soporte y pruebas\n\n'
support = [
('src/environments/environment.ts','Extiende sharedEnvironment con production:false.'),
('src/environments/environment.production.ts','Extiende sharedEnvironment con production:true.'),
('src/app/pages/login/login.component.css','Referencia visual: fondo oscuro y tonos oliva; no se modifica.'),
('src/app/pages/sidebar/sidebar.css','Diseño de navegación; no se modifica.'),
('src/app/app.ts','Raíz con RouterOutlet; no instancia NavigationService.'),
('src/app/core/auth/auth-test.fixture.ts','Bootstrap sintético para pruebas, sin credenciales reales; no es respuesta de producción.'),
('src/app/core/auth/login-contract.spec.ts','Usuario, permisos textuales, cinco catálogos, jerarquía, errores y atomicidad.'),
('src/app/core/auth/logout.spec.ts','Cierre local, cancelación, aislamiento entre sesiones y limpieza.'),
('src/app/core/auth/authorization-policy.spec.ts','Matriz numérica aislada e IDs inválidos.'),
('src/app/core/services/api-connection.spec.ts','URL-encoded, expiración, guard, destinos HTTP y respuestas 401/403.'),
('src/app/core/interceptors/loader.interceptor.spec.ts','Liberación de operaciones del loader, también en error/cancelación.'),
('src/app/pages/login/login.component.spec.ts','Envío único, navegación y errores seguros del formulario.'),
('src/app/pages/sidebar/sidebar.spec.ts','Renderizado y acciones de salida de escritorio/móvil.'),
('src/app/layout/app-shell.spec.ts','Vistas protegidas y limpieza de datos al cambiar usuario.'),
('angular.json','Builders, proxy de serve y presupuestos de compilación.'),
('package.json','Dependencias y comandos npm; versiones instaladas en contexto-revision.json.'),
('src/app/core/auth/README.md','Material histórico; contrastar con código y este informe antes de reutilizar.'),
('docs/login-contrato-fastapi.md','Informe anterior de adaptación; no sustituye la ejecución de pruebas de este corte.'),
('docs/autenticacion-v1-pendientes.md','Preparación de política numérica y contratos v1 pendientes.'),
]
invmd += mdtable(['Ruta','Responsabilidad'],support)
(BASE / 'INVENTARIO_ARCHIVOS_AUTENTICACION.md').write_text(invmd,encoding='utf-8')

bitrows=[]
for item in META['historial']:
    commit,date,msg=item.split('|',2)
    bitrows.append((date, f'Commit {commit}', 'Árbol Git', msg, 'Motivo específico no inferido del asunto', 'Commit existente; no acredita sprint terminado', 'git log; contexto-revision.json', 'Relacionar requerimientos con commits propios'))
bitrows += [
('2026-10-27','Revisión técnica del árbol actual','Archivos F-01 a F-24','Inventario de mecanismos reales y brechas v1','Sustentar el sprint con código, sin presumir backend','Inventario y riesgos registrados','INVENTARIO_ARCHIVOS_AUTENTICACION.md','Confirmar OpenAPI v1 y seed'),
('2026-10-27','Documentación TSDoc','15 archivos TS listados en verificacion-solo-comentarios.json','Comentarios en español sobre límites, entradas, retornos y seguridad','Transferencia de conocimiento sin cambio de firmas/lógica','JavaScript idéntico en 15/15 archivos','evidencias/verificacion-solo-comentarios.json','Revisión editorial por responsable'),
('2026-10-27','Verificación Angular','Pruebas y builders actuales','Selección 53/53, global 166 aprobadas y 10 fallidas; desarrollo correcto; producción con presupuesto CSS excedido','Registrar ejecución efectiva','Resultados reproducibles guardados','evidencias/*.log','Atender fallos globales en alcance separado'),
('2026-10-27','Generación documental','Cinco entregables y evidencias','Word, arquitectura, matriz, bitácora e inventario','Sustento institucional trazable','Archivos generados desde fuentes locales','generar_documentacion.py; evidencias','Completar responsable/estado y revision editorial'),
]
bitmd='# Bitácora técnica de autenticación\n\nNo se atribuyen autores ni fechas de implementación a cambios no confirmados por Git. Las entradas del 27/10/2026 corresponden a revisión/documentación y ejecución de este corte, no a la fecha original de desarrollo de cada funcionalidad. El único commit encontrado en las rutas auth/login/interceptores es b33a1d2; otros cambios del árbol están sin commit.\n\n'+mdtable(['Fecha','Actividad','Archivo','Cambio','Motivo','Resultado','Evidencia','Pendiente'],bitrows)
(BASE/'BITACORA_DESARROLLO_AUTENTICACION.md').write_text(bitmd,encoding='utf-8')

def font(size,bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'),size)
def wrap(text,width=27):
    import textwrap
    return '\n'.join(textwrap.wrap(text,width=width))
def diagram_arch():
    im=Image.new('RGB',(1600,960),'white'); d=ImageDraw.Draw(im)
    def box(x,y,w,h,title,body,muted=False):
        d.rounded_rectangle((x,y,x+w,y+h),radius=18,fill='#F3F4EE',outline='#78813E' if not muted else '#8C8C8C',width=4)
        d.text((x+18,y+16),title,font=font(30,True),fill='#172419')
        d.multiline_text((x+18,y+65),wrap(body,25),font=font(28),fill='#313A31',spacing=7)
    def arrow(a,b,label=None):
        d.line([a,b],fill='#78813E',width=5)
        import math
        ang=math.atan2(b[1]-a[1],b[0]-a[0]); tip=[b,(b[0]-17*__import__('math').cos(ang-.5),b[1]-17*__import__('math').sin(ang-.5)),(b[0]-17*__import__('math').cos(ang+.5),b[1]-17*__import__('math').sin(ang+.5))]
        d.polygon(tip,fill='#78813E')
        if label:d.text(((a[0]+b[0])//2-45,(a[1]+b[1])//2-40),label,font=font(24),fill='#172419')
    d.text((35,20),'ARQUITECTURA OBSERVADA | autenticación actual',font=font(36,True),fill='#172419')
    box(35,100,460,215,'Presentación Angular','LoginComponent, Sidebar, Router y AppShell')
    box(570,100,460,215,'AuthService','Signals, generación local y validación bootstrap')
    box(1105,100,460,215,'Datos de sesión','Usuario, tokens en memoria y CatalogosService')
    arrow((495,205),(570,205));arrow((1030,205),(1105,205))
    box(35,410,460,215,'Guards e interceptores','Autenticación local, Bearer limitado y limpieza')
    box(570,410,460,215,'HttpClient / proxy','/api/auth/token se traduce a POST /token')
    box(1105,410,460,215,'FastAPI local','OpenAPI consultado. Sin prueba de login real')
    arrow((800,315),(800,410));arrow((570,515),(495,515));arrow((1030,515),(1105,515))
    box(35,720,730,170,'Preparado, sin integración','Política numérica: reglas aisladas y probadas',True)
    box(830,720,735,170,'No verificable / pendiente','JWT real, BD, Redis, refresh, /yo y seguridad v1',True)
    im.save(BASE/'arquitectura.png')
def diagram_seq():
    im=Image.new('RGB',(1600,1120),'white');d=ImageDraw.Draw(im)
    xs=[160,570,1000,1430];labels=['Usuario','Login Angular','AuthService','FastAPI / proxy']
    for x,label in zip(xs,labels):
        d.rounded_rectangle((x-145,25,x+145,100),radius=10,fill='#172419');d.text((x-130,45),label,font=font(27,True),fill='white')
        d.line((x,105,x,1060),fill='#B8BEAD',width=3)
    def msg(a,b,y,label,ret=False):
        d.line((xs[a],y,xs[b],y),fill='#78813E',width=4)
        sign=1 if b>a else -1
        d.polygon([(xs[b],y),(xs[b]-sign*15,y-8),(xs[b]-sign*15,y+8)],fill='#78813E')
        d.text((min(xs[a],xs[b])+12,y-43),label,font=font(26),fill='#172419')
    msg(0,1,180,'Formulario válido')
    msg(1,2,290,'login(usuario,password)')
    msg(2,3,405,'POST /token URL-encoded')
    d.rounded_rectangle((1120,435,1580,545),radius=10,fill='#F3F4EE',outline='#888888',width=2)
    d.multiline_text((1135,448),'Credenciales / firma JWT:\ninternos no auditados',font=font(26),fill='#313A31',spacing=9)
    msg(3,2,620,'Bootstrap esperado')
    d.rounded_rectangle((810,660,1195,785),radius=10,fill='#F3F4EE',outline='#78813E',width=3)
    d.multiline_text((830,675),'Validar todo; copiar\ncatálogos y publicar sesión',font=font(26),fill='#172419',spacing=9)
    msg(2,1,855,'next o error controlado')
    msg(1,0,965,'Home o mensaje seguro')
    d.text((45,1070),'No hay refresh, /yo, rotación o logout remoto en este flujo.',font=font(29,True),fill='#172419')
    im.save(BASE/'secuencia.png')
diagram_arch();diagram_seq()

doc=Document();sec=doc.sections[0]
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
styles=doc.styles
for name in ['Normal','Body Text']:
    styles[name].font.name='Calibri';styles[name].font.size=Pt(10)
    styles[name].paragraph_format.space_after=Pt(6)
for name,size in [('Title',30),('Heading 1',18),('Heading 2',12),('Heading 3',10.5)]:
    styles[name].font.name='Calibri';styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor.from_string(DARK)
    styles[name].paragraph_format.keep_with_next=True
for name in ['Subtitle','Caption']:
    styles[name].font.color.rgb=RGBColor.from_string(OLIVE)
for part in doc.part.package.parts:
    if str(part.partname) == '/word/theme/theme1.xml':
        theme = etree.fromstring(part.blob)
        ns = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for key, value in [('accent1',OLIVE),('accent2',DARK),('hlink',OLIVE),('folHlink',DARK)]:
            node = theme.find('.//a:clrScheme/a:'+key,ns)
            for child in list(node):node.remove(child)
            node.append(etree.Element('{'+ns['a']+'}srgbClr',val=value))
        part._blob = etree.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
header=sec.header.paragraphs[0];header.text='VULCANO  |  DOCUMENTACIÓN TÉCNICA  |  AUTENTICACIÓN';header.style='Caption'
header.runs[0].font.color.rgb=RGBColor.from_string(OLIVE)
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('Versión 1.0  •  27/10/2026  |  Página ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
footer.add_run(' de ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'NUMPAGES');footer._p.append(fld)
doc.core_properties.title='VULCANO — Sustento técnico del sprint de autenticación y seguridad'
doc.core_properties.subject='Estado del código, evidencias de pruebas, arquitectura y actividades pendientes'
doc.core_properties.author='Responsable por definir'
doc.core_properties.keywords='VULCANO, Angular, autenticación, sprint, evidencia'
def p(text,style=None):return doc.add_paragraph(text,style)
def h(text,l=1):return doc.add_heading(text,l)
def table(headers,rows):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(4)
                for run in paragraph.runs:run.font.size=Pt(8 if len(headers)>4 else 9)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),DARK);cell._tc.get_or_add_tcPr().append(shading)
                for run in cell.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
    p('')
    return t
def newsection(title):doc.add_page_break();h(title)

p('VULCANO', 'Title');p('SUSTENTO TÉCNICO DEL SPRINT\nDE AUTENTICACIÓN Y SEGURIDAD','Subtitle')
p('Módulo: Login y administración de sesiones')
table(['Control documental','Valor'],[('Versión','1.0'),('Fecha de elaboración',DATE),('Responsable','[Campo editable: nombre y cargo]'),('Estado','[Campo editable: revisión / aprobación]'),('Corte técnico','Árbol de trabajo inspeccionado el 27/10/2026'),('Alcance de evidencia','Código frontend, OpenAPI local y pruebas automatizadas; sin auditoría de backend o prueba de login real')])
p('Documento de transferencia técnica y sustento del sprint. Las capacidades pendientes se identifican expresamente; no constituye certificación de seguridad o autorización de despliegue.')
p('Identidad visual documental: fondo claro, encabezados oscuros y acento oliva derivados de la interfaz existente. No se modifica el diseño de la aplicación.')
newsection('Contenido y lectura de estados')
for title in ['1. Resumen ejecutivo','2. Antecedentes y justificación','3. Objetivo general y objetivos específicos','4. Alcance del sprint','5. Arquitectura técnica','6. Desarrollo del módulo de login','7. Implementación de seguridad','8. Integración frontend–backend','9. Documentación del código fuente','10. Historias de usuario y criterios de aceptación','11. Pruebas y validación','12. Resultados y avances del sprint','13. Riesgos y actividades pendientes','14. Conclusiones','15. Anexos']:
    p(title)
table(['Estado','Criterio de uso'],[(OK,'Código existente y caso ejecutado aplicable; indicar cuando se usa backend simulado.'),('Implementado sin validar','Código/entregable presente, sin ejecución de su comprobación correspondiente.'),(PART,'Parte del requisito funciona; la integración o validación restante se identifica.'),(PEND,'Sin implementación del requisito en el frontend revisado.'),(NV,'Falta evidencia del servidor o ejecución real; no se infiere a partir del requerimiento.')])
newsection('1. Resumen ejecutivo')
p('El sprint consolidó el login existente de VULCANO mediante un servicio de sesión centralizado, interfaces de bootstrap, validación previa a la autenticación, catálogos de sesión y cierre local con aislamiento de respuestas y vistas privadas. Las verificaciones actuales confirman 53 pruebas enfocadas aprobadas y compilación de desarrollo correcta. La suite completa presenta 166 pruebas aprobadas y 10 fallidas; producción no supera el presupuesto de estilos de Hoja Milimetrada.')
p('La plataforma aún no integra los endpoints /api/v1/auth/*. El código utiliza POST /api/auth/token traducido a /token por el proxy de desarrollo. El OpenAPI local consultado publica /token pero no las seis rutas nuevas. No se efectuaron solicitudes de login con credenciales reales, renovación, revocación o desbloqueo. La política de permisos numéricos existe como funciones puras probadas, sin conexión a menús, rutas o acciones.')
p('El estado global del sprint es parcialmente implementado. Refresh rotativo, sesión deslizante, recuperación /yo, revocación Redis y administración de IPs continúan pendientes de contrato e implementación. Este informe distingue los avances verificables de esos objetivos para evitar dar por terminadas capacidades que todavía no operan.')
h('2. Antecedentes y justificación')
p('El contrato de autenticación evolucionó desde campos opcionales antiguos hacia una respuesta bootstrap con usuario y catálogos. Aceptar solo un token sin validar el resto podía dejar datos incompletos para la navegación. Centralizar validación y sesión reduce ese riesgo y evita que varios componentes creen copias inconsistentes de identidad.')
p('Los riesgos abordados son uso local de credenciales vencidas, reconstrucción de sesión por una respuesta tardía después de salir, reutilización de datos privados entre usuarios y exposición del token a destinos externos. La autorización detallada por función, revocación remota y prevención de fuerza bruta requieren aún coordinación con el contrato/backend; no se presume su cobertura.')
h('3. Objetivo general y objetivos específicos')
p('Sustentar técnicamente el estado del módulo de login y administración de sesiones, con trazabilidad desde requisito hasta código y verificación ejecutada.')
for t in ['Identificar archivos, contratos, dependencias y consumidores de identidad.', 'Explicar el flujo real del bootstrap, expiración y cierre local.', 'Separar autenticación operativa de la política numérica preparada y de la arquitectura v1 futura.', 'Registrar resultados de pruebas/builds y hallazgos priorizados sin cambiar comportamiento.', 'Entregar Word y anexos mantenibles para transferencia del conocimiento.']:p(t,'List Bullet')
h('4. Alcance del sprint')
table(['Incluido y observado','Excluido / pendiente'],[('Formulario existente, URL-encoded, mensajes seguros y validación integral.','Sin reconstrucción del login ni cambio visual.'),('Estado en memoria, expiración por expira_en y catálogos.','Sin persistencia de refresh ni recuperación /yo.'),('Guards de autenticación, límites de Bearer y logout local.','Sin revocación remota ni validación criptográfica Angular.'),('Política numérica probada en aislamiento.','Sin IDs asociados a funciones ni aplicación en menús/rutas/acciones.'),('Comentarios TSDoc, documentación y pruebas.','Sin cambios funcionales a mapas, cálculos, backend o herramientas operacionales.')])
newsection('5. Arquitectura técnica')
p('La arquitectura usa componentes standalone Angular, servicios providedIn:root, signals y RxJS. AuthService centraliza tokens, usuario, expiración y generación de sesión; CatalogosService mantiene el bootstrap sin HTTP adicional. LoginComponent y Sidebar consumen esa fuente. AppShell registra limpieza de herramientas privadas y mantiene PuntosDemoService por instancia.')
table(['Tecnología instalada','Versión comprobada'],META['versiones_instaladas'].items())
doc.add_picture(str(BASE/'arquitectura.png'),width=Inches(6.6));p('Figura 1. Arquitectura observada. Los bloques inferiores son preparación o aspectos no verificables; no se representan como integración operativa.','Caption')
p('FastAPI se acredita por el OpenAPI local de “Servidor VULCANO”, versión 2.0.0. No se inspeccionó su código fuente de autenticación ni la infraestructura. La generación/verificación JWT, base de datos de usuarios/roles/permisos y Redis se describen como responsabilidades esperadas del servidor, no como implementación comprobada.')
doc.add_picture(str(BASE/'secuencia.png'),width=Inches(6.6));p('Figura 2. Secuencia real esperada por el frontend y simulada en pruebas. El contrato HTTP está publicado; el éxito autenticado real no fue ejecutado.','Caption')
h('5.1 Límites de confianza',2)
p('El JSON HTTP no se considera válido por el tipo TypeScript: validarLogin comprueba todos los campos antes de publicar la sesión. Un token presente en memoria no demuestra firma ni autorización remota. El servidor debe verificar autenticidad, revocación y autorización. El frontend no contiene clave secreta JWT ni intenta verificar con ella.')
newsection('6. Desarrollo del módulo de login')
for step in [
 'El formulario requiere usuario y contraseña. El envío se detiene si ya existe una petición pendiente o controles inválidos. No se implementa un segundo login.',
 'AuthService.login devuelve un Observable diferido, cancela el intento previo y envía username/password URL-encoded a ApiConfigService.loginUrl. HttpContext excluye Bearer previo y el logout global por error de credenciales.',
 'El timeout configurado es 30000 ms. validarLogin exige access_token/refresh_token no vacíos, tipo_token bearer, expira_en positivo y finito, usuario completo y cinco catálogos con registros compatibles.',
 'El servicio copia usuario y catálogos; calcula expiresAt con Date.now()+expira_en*1000. Publica el access token tras cargar catálogos. No usa token_type/expires_in, no decodifica claims y no consulta perfil adicional.',
 'El componente navega a /home solo al recibir next válido. finalize libera el estado de envío; takeUntilDestroyed cancela al destruir la vista. Los errores no exponen contenido arbitrario del servidor.',
 'Logout incrementa la generación, limpia estado/catálogos, cancela login pendiente, retira claves propias antiguas y ejecuta callbacks privados antes de redirigir a /login con replaceUrl. La promesa concurrente es compartida. No se comunica con el servidor.'
]:p(step,'List Number')
h('6.1 Bootstrap y catálogos',2)
table(['Campo / conjunto','Tratamiento actual'],[('access_token / refresh_token','Memoria; solo access se usa en Authorization. Refresh reservado para integración futura.'),('tipo_token / expira_en','Se exige bearer; expira_en determina vencimiento local.'),('usuario','Identidad, grado, correo, rol, permisos textuales, aprobación de documentos y foto opcional.'),('tipos_unidad / unidades','Almacenamiento en CatalogosService; jerarquía mediante unidad_padre_id.'),('estados_operacionales / estados_logisticos / permisos','Catálogos disponibles sin nuevas consultas; permisos del catálogo no son asignaciones.')])
p('La fotografía relativa se resuelve contra el origen auth configurado y solo admite HTTP(S). Sin foto retorna null. Sidebar muestra grado y nombre completo. El código preparado de NavigationService no se encontró instanciado, por lo que no se afirma trazabilidad activa de navegación.')
newsection('7. Implementación de seguridad')
h('7.1 Tokens y claims',2)
p('El frontend trata los tokens como cadenas opacas y no implementa JwtPayload ni lee claims. sub, rol, permisos:number[], exp, iat y jti son claims requeridos para la futura integración, pendientes de comprobar en un JWT real anonimizado. expira_en del bootstrap es la única base actual de expiración. No se afirma que los tokens emitidos realmente contengan esos claims o que una firma se haya validado en esta revisión.')
table(['Claim requerido','Estado de evidencia frontend'],[(c,'Sin lectura/validación implementada; confirmar contrato real') for c in ['sub','rol: ADMINISTRADOR | USUARIO','permisos: number[]','exp','iat (opcional)','jti (opcional)']])
h('7.2 Matriz de roles y permisos',2)
table(['Rol / asignaciones','Política numérica preparada','Estado operacional'],[('ADMINISTRADOR / []','Permite cualquier ID válido','No integrada. AuthService textual no concede excepción.'),('ADMINISTRADOR / lista','Solo IDs de la lista','No integrada a funciones.'),('USUARIO / []','No concede acceso; conserva autenticado el contexto','Sin guard de función ni alerta específica implementada.'),('USUARIO / lista','Solo IDs de la lista','Pendientes seed y enlaces menú/ruta/acción.'),('Sin sesión / rol desconocido','Deniega','authGuard sí exige sesión, pero no aplica la política numérica.')])
p('Los IDs de catálogo no equivalen a permisos asignados. AuthService.tienePermiso(codigo:string) verifica una lista textual del usuario, no todos los registros del catálogo. Sidebar.canShowItem usa un booleano estático; permisoCodigo no se utiliza para autorizar. Las funciones puras tienen cobertura aislada y no deben presentarse como control operacional completo.')
h('7.3 Sesión deslizante y renovación',2)
p('Access token de 25 minutos e inactividad de 60 minutos son requisitos nuevos. No existe contador de actividad humana, renovación periódica, refresh compartido, rotación ni reintento de 401. El código solo comprueba expiración al consultar autenticación/token o ejecutar guards. No hay renovación infinita por tráfico automático porque no hay renovación implementada. El siguiente sprint debe medir interacciones reales, no telemetría o teselas, y respetar el vencimiento impuesto por el servidor.')
h('7.4 Persistencia, cierre y revocación',2)
p('Tokens y usuario están en memoria, reduciendo su permanencia en almacenamiento del navegador sin eliminar el riesgo de XSS. Recargar implica nuevo login. Se retiran claves históricas propias de ambos almacenamientos; otras preferencias permanecen. Cookies HttpOnly no pueden limpiarse desde este código. Logout no confirma revocación de access/refresh en Redis; tampoco se consultó Redis o una base de datos. No hay mecanismo comprobado de lista de revocación o control de fuerza bruta del servidor.')
h('7.5 Guards, interceptores y errores',2)
p('authGuard protege rutas del shell y guestGuard evita abrir login con sesión activa. authInterceptor adjunta access solo a bases autorizadas y filtra respuestas exitosas obsoletas. errorInterceptor conserva sesión ante 403, cierra localmente ante 401 protegido actual y propaga errores. Los 5xx se registran completos: se identifica oportunidad de redacción segura de logs. No existe guard administrador, panel de bloqueos o autorización por IDs en acciones internas.')
newsection('8. Integración frontend–backend')
p('Métodos y esquemas se documentan solo cuando tienen evidencia. La existencia de una URL configurada no prueba un consumidor implementado. El OpenAPI consultado no enumera /api/v1/auth/*; su ausencia en ese documento no demuestra que no exista otro servidor o versión. No se probaron credenciales ni operaciones administrativas reales.')
for ep,method,io,auth,consumer,codes,status,evidence in endpoints:
    h(ep,2)
    table(['Propiedad','Evidencia actual'],[('Método real',method),('Entrada / salida',io),('Autenticación',auth),('Consumidor',consumer),('Respuestas y errores',codes),('Estado',status),('Sustento',evidence)])
p('Configuradas pero no consumidas: /api/auth/logout y /api/auth/refresh. No deben confundirse con las rutas v1. La semántica de seguridad requerida por el usuario (revocación Redis, bloqueo temporal, intentos fallidos y desbloqueo manual) es una dependencia de contrato, no una función demostrada por el OpenAPI consultado.')
newsection('9. Documentación del código fuente')
p('Se añadieron comentarios y TSDoc en español sin alterar firmas o instrucciones ejecutables. La comparación SHA-256 del JavaScript transpilado con removeComments=true confirma igualdad en 15/15 archivos. Esta comprobación complementa, pero no sustituye, las pruebas del comportamiento.')
table(['Grupo','Archivos comentados en este corte'],[('Autenticación','auth.service.ts; auth-response.validator.ts; catalogos.service.ts; authorization-policy.ts; auth.guard.ts; guest.guard.ts'),('HTTP y servicios','auth.interceptor.ts; error.interceptor.ts; api-config.service.ts; navigation.service.ts'),('Consumidores','login.component.ts; sidebar.ts; app-shell.ts')])
p('El conjunto de comparación también incluye auth.models.ts y auth-http-context.ts, revisados sin necesitar comentarios adicionales. Los cambios anteriores del árbol ya estaban presentes al iniciar esta documentación: no se atribuyen a esta fase ni a un autor sin evidencia.')
p('El inventario detallado de funciones, parámetros, retornos, errores y dependencias se incorpora en el anexo 15.1 y en INVENTARIO_ARCHIVOS_AUTENTICACION.md. Las firmas se extraen del AST; interfaces, plantillas y configuración no se presentan como métodos ejecutables.')
newsection('10. Historias de usuario y criterios de aceptación')
for id_,need,criteria,status,evidence in histories:
    h(id_+' — '+need,2);p('Criterios de aceptación: '+criteria);p('Estado: '+status);p('Evidencia y límites: '+evidence)
h('10.1 Matriz de trazabilidad',2)
table(['Historia','Requisito','Archivo','Endpoint','Prueba','Estado'],trace)
newsection('11. Pruebas y validación')
p('Pruebas automatizadas ejecutadas durante este corte, con HttpTestingController y entorno de pruebas Angular/Vitest. Las pruebas de componentes y RouterTestingHarness validan integración interna con respuestas simuladas; no son una prueba funcional de FastAPI/Redis ni una validación manual en producción. No se registran tokens reales ni credenciales en los entregables.')
table(['ID','Caso','Esperado','Obtenido','Estado','Evidencia'],tests)
h('11.1 Comandos y resultados',2)
for command in ['npm test -- --watch=false','npm test -- --watch=false --include=src/app/core/auth/*.spec.ts --include=src/app/core/services/api-connection.spec.ts --include=src/app/core/interceptors/*.spec.ts --include=src/app/pages/login/login.component.spec.ts --include=src/app/pages/sidebar/sidebar.spec.ts --include=src/app/layout/app-shell.spec.ts','npx ng build --configuration development','npx ng build']:
    p(command)
p('Resultado enfocado: 8 archivos y 53 pruebas aprobadas. Global: 37 archivos, 176 pruebas; 5 archivos con 10 fallos y 32 archivos aprobados con 166 casos. Desarrollo compila. Producción falla por hoja-milimetrada.component.css (19,26 kB frente a 18 kB) y registra avisos CommonJS de Leaflet/MapLibre. No se modifica ese módulo ni los budgets.')
h('11.2 Fallos de la suite global',2)
table(['Archivo','Resultado observado'],[('src/app/app.spec.ts','1 fallo: expectativa del título inicial.'),('src/app/pages/vulcano/menu-sistemas/menu-sistemas.spec.ts','2 fallos: IDs de trabajos esperados distintos del estado actual.'),('src/app/pages/vulcano/ventana-sistemas-artilleria/ventana-sistemas-artilleria.spec.ts','2 fallos: destinos/contenido provisional esperados.'),('src/app/pages/vulcano/hoja-milimetrada/proyectar.spec.ts','1 fallo: registrarCoordenadas no fue invocado según expectativa.'),('src/app/shared/components/mapa/mapa.spec.ts','4 fallos: casos de trayectoria/marcadores en Leaflet; ver log para pila completa.')])
p('Estos archivos no se modificaron en la fase documental. Los fallos se reportan tal como se ejecutaron; no se atribuyen a una causa única ni se ocultan con cambios de expectativas. La igualdad de JavaScript en el código comentado permite descartar alteraciones ejecutables introducidas por TSDoc.')
newsection('12. Resultados y avances del sprint')
table(['Avance comprobado','Contribución / límite'],[('Bootstrap estricto y centralizado','Impide publicar una sesión con usuario o catálogos incompletos.'),('Tokens solo en memoria y expiración local','Evita restauración automática de tokens persistidos antiguos; no equivale a revocación remota.'),('Generación local y cancelación de login','Reduce repoblación de estado privado por respuestas tardías.'),('Limpieza del shell y catálogos','Aísla vistas y datos entre sesiones.'),('Bearer limitado por destino','Evita enviar credenciales a dominios de mapas u otros orígenes no configurados.'),('Política numérica probada','Prepara reglas del contrato nuevo; conexión al producto pendiente.'),('Documentación institucional y TSDoc','Explica entradas, retornos y riesgos sin variar comportamiento.')])
p('El comparativo con el contrato anterior se sustenta en los informes locales y el árbol disponible. El historial Git no permite fechar cada mejora por separado; no se cuantifica productividad, velocidad del sprint o esfuerzo sin registros. El corte demuestra capacidades existentes y documentación generada, no la finalización de la arquitectura de seguridad v1.')
h('12.1 Bitácora y evidencia Git',2)
p('Commits disponibles: '+ '; '.join(META['historial'])+'. El historial sobre rutas auth/login/interceptores contiene b33a1d2. Hay cambios sin commit en el árbol; no se inventan autores o fechas de implementación. La bitácora distingue el historial de repositorio de las actividades de revisión realizadas hoy.')
table(['Fecha','Actividad','Resultado / evidencia'],[(r[0],r[1],r[5]+' — '+r[6]) for r in bitrows])
newsection('13. Riesgos y actividades pendientes')
for id_,priority,finding,impact,reco in risks:
    h(id_+' | Prioridad '+priority,2);p('Hallazgo: '+finding);p('Impacto o limitación: '+impact);p('Recomendación: '+reco)
h('13.1 Plan del siguiente sprint',2)
for t in ['Confirmar OpenAPI v1, ejemplos anónimos de bootstrap/JWT, seed de permisos y configuración de proxy/HTTPS.', 'Adaptar el contrato numérico en el servicio central y conectar cada función a un ID documentado con controles de ruta y acción.', 'Implementar refresh rotativo compartido, un reintento máximo y actividad humana eficiente; pruebas de concurrencia e inactividad.', 'Integrar logout remoto con resultado explícito de revocación; definir recuperación /yo y catálogos tras recarga sin endpoints inventados.', 'Crear seguridad administrativa con guard de rol, confirmación de IP y actualización tras desbloqueo real.', 'Resolver fallos globales/build y probar contrato real con backend autorizado antes de declarar la integración terminada.']:p(t,'List Number')
h('14. Conclusiones')
p('VULCANO dispone de una base frontend comprobable para autenticación bootstrap y cierre local: estado central, validación integral, catálogos, expiración y limpieza privada. La documentación identifica sus componentes y aporta evidencias reproducibles de 53 casos enfocados aprobados. El módulo aún no cumple de extremo a extremo los requisitos de permisos numéricos operativos, sesión deslizante o administración Redis. Esas capacidades exigen contrato y desarrollo posterior; su ausencia no debe confundirse con una integración terminada.')
p('La transferencia queda sustentada mediante Word, diagramas, inventario, trazabilidad y bitácora. Los comentarios conservan el JavaScript generado y no introducen cambios funcionales. La aceptación institucional requiere completar los campos de responsable/estado y revisar los hallazgos y resultados globales antes de planear el despliegue.')
newsection('15. Anexos')
h('15.1 Matriz de archivos y contratos públicos',2)
for i,row in enumerate(inventory,1):
    file,name,role,deps,ep,security,err,obs=row
    h(f'F-{i:02d}. {name}',3);p('Ruta: '+file)
    for label,value in [('Responsabilidad',role),('Dependencias y consumidores',deps),('Endpoints',ep),('Seguridad',security),('Errores',err),('Observaciones',obs)]:p(label+': '+value)
    methods=[s for s in SYMBOLS if s['archivo']==file]
    if methods:
        table(['Método','Entradas','Retorno'],[(s['nombre'],s['entrada'] or 'Sin parámetros',s['retorno']) for s in methods])
    else:p('No se extraen métodos de clase de este archivo. Entradas y retornos de guards/interceptores se describen en su responsabilidad; las plantillas y configuraciones no son servicios HTTP.')
h('15.2 Soporte, configuración y pruebas',2);table(['Ruta','Responsabilidad / alcance'],support)
h('15.3 Fragmentos relevantes sin credenciales',2)
p('Extractos de AuthService y política preparada, incluidos para aclarar la diferencia de contratos:')
for snippet in ["obtenerPermisos(): string[] { return [...(this.currentUser()?.permisos ?? [])]; }", "tienePermiso(codigo: string): boolean { return this.isAuthenticated() && this.obtenerPermisos().includes(codigo); }", "return (contexto.rol === 'ADMINISTRADOR' && contexto.permisos.length === 0) || contexto.permisos.includes(id);", "expiresAt: Date.now() + response.expira_en * 1000"]:
    run=p(snippet).runs[0];run.font.name='Consolas';run.font.size=Pt(8)
h('15.4 Evidencias y referencias locales',2)
refs = ['ARQUITECTURA_AUTENTICACION.md','MATRIZ_TRAZABILIDAD_AUTENTICACION.md','BITACORA_DESARROLLO_AUTENTICACION.md','INVENTARIO_ARCHIVOS_AUTENTICACION.md','evidencias/contexto-revision.json','evidencias/inventario-simbolos.json','evidencias/verificacion-solo-comentarios.json','evidencias/pruebas-autenticacion.log','evidencias/pruebas-completas.log','evidencias/compilacion-desarrollo.log','evidencias/compilacion-produccion.log','evidencias/validacion-documento.json']
p('; '.join(refs))
p('Fuentes primarias: código frontend revisado, árbol/historial Git y OpenAPI local consultado sin credenciales. Las especificaciones de Redis/JWT v1 son requerimientos recibidos y no una auditoría del backend. No se incorporan datos personales de usuarios reales ni secretos.')
doc.save(BASE/'VULCANO_SPRINT_AUTENTICACION_SEGURIDAD.docx')

manifest={}
for row in inventory:
    path=ROOT/row[0];manifest[row[0]]=hashlib.sha256(path.read_bytes()).hexdigest()
(EV/'huellas-fuentes.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print('Generados Word, cuatro anexos Markdown y diagramas a partir de evidencia local.')
